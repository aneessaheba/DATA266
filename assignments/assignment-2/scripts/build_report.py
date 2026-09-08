"""Builds the Assignment 2 Word report with write-up text, real result tables,
and clearly marked screenshot placeholders for manual insertion."""
import docx
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = "/Users/anees/Documents/Coursework/data266-5330/assignments/assignment-2/HW2_Document.docx"

doc = docx.Document()


def set_cell_shading(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tcPr.append(shd)


def add_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "999999")
        borders.append(el)
    tblPr.append(borders)


def add_table(rows, header=True):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    add_table_borders(table)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.rows[i].cells[j]
            cell.text = str(val)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9.5)
                    if header and i == 0:
                        r.font.bold = True
            if header and i == 0:
                set_cell_shading(cell, "D9D9D9")
    doc.add_paragraph()
    return table


def add_screenshot_placeholder(description):
    table = doc.add_table(rows=1, cols=1)
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "dashed")
        el.set(qn("w:sz"), "8")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "888888")
        borders.append(el)
    tblPr.append(borders)
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, "F2F2F2")
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(18)
    run = p.add_run(f"[SCREENSHOT PLACEHOLDER]\n{description}")
    run.italic = True
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    run.font.size = Pt(10)
    doc.add_paragraph()


def h1(text):
    doc.add_heading(text, level=1)


def h2(text):
    doc.add_heading(text, level=2)


def p(text):
    doc.add_paragraph(text)


# ============================================================ Title
doc.add_heading("DATA 266 HW2", level=0)
p("Embedding-based transfer learning on IMDB reviews, a Retrieval-Augmented Generation "
  "pipeline over 10 movie Wikipedia pages, and five training-time optimization techniques.")

# ============================================================ Part 1
h1("1. Embedding-Based Transfer Learning (IMDB Reviews)")

h2("1.1 Setup")
p("Pretrained model: word2vec-google-news-300 (3,000,000 words, 300 dimensions), loaded with "
  "gensim's KeyedVectors. Fine-tuning corpus: 15,000 reviews randomly sampled from the IMDB "
  "Movie Reviews dataset (stanfordnlp/imdb, train + unsupervised splits, 75,000 reviews "
  "available total), cleaned of HTML line breaks and tokenized with gensim's simple_preprocess. "
  "A new gensim Word2Vec model (skip-gram, 300 dimensions) was built over the IMDB vocabulary, "
  "every word shared with the pretrained vocabulary was initialized with its pretrained vector "
  "(20,040 of 22,723 IMDB-vocab words matched), and training continued for 5 epochs on the IMDB "
  "corpus.")

h2("1.2 Top-3 nearest neighbors before fine-tuning")
add_table([
    ["Word", "Rank", "Neighbor", "Cosine similarity"],
    ["cast", 1, "casts", 0.7219],
    ["cast", 2, "casting", 0.7188],
    ["cast", 3, "Cast", 0.6638],
    ["score", 1, "scoring", 0.7197],
    ["score", 2, "scores", 0.6596],
    ["score", 3, "scored", 0.6384],
    ["plot", 1, "plots", 0.7625],
    ["plot", 2, "Plot", 0.6524],
    ["plot", 3, "plotting", 0.6328],
    ["screen", 1, "screens", 0.7729],
    ["screen", 2, "onscreen", 0.6115],
    ["screen", 3, "LCD_screen", 0.5599],
    ["review", 1, "reviewed", 0.6630],
    ["review", 2, "reviewing", 0.6610],
    ["review", 3, "reviews", 0.6380],
])
add_screenshot_placeholder("Notebook cell output showing the pretrained-model neighbor table "
                            "(Part 1, Section 1.2 of the notebook).")

h2("1.3 Top-3 nearest neighbors after fine-tuning")
add_table([
    ["Word", "Rank", "Neighbor", "Cosine similarity"],
    ["cast", 1, "supporting", 0.5975],
    ["cast", 2, "ensemble", 0.5629],
    ["cast", 3, "casts", 0.5266],
    ["score", 1, "ennio", 0.5714],
    ["score", 2, "morricone", 0.5708],
    ["score", 3, "scoring", 0.5438],
    ["plot", 1, "story", 0.6103],
    ["plot", 2, "storyline", 0.5662],
    ["plot", 3, "plotline", 0.5534],
    ["screen", 1, "screens", 0.5327],
    ["screen", 2, "onscreen", 0.5216],
    ["screen", 3, "kapadia", 0.3945],
    ["review", 1, "reviews", 0.6052],
    ["review", 2, "comment", 0.5809],
    ["review", 3, "comments", 0.5578],
])
p("Every target word's neighbors shift toward its domain-specific, film-critique sense after "
  "fine-tuning: cast toward supporting/ensemble (an ensemble cast of actors), score toward ennio/"
  "morricone (the film composer Ennio Morricone), plot toward story/storyline (narrative plot, "
  "not a literal plot of land), and review toward reviews/comment/comments (review-discussion "
  "context). screen was already film-adjacent pre-fine-tuning and shifts the least in sense.")
add_screenshot_placeholder("Notebook cell output showing the fine-tuned-model neighbor table, "
                            "and the fine-tuning training log (IMDB load, tokenization, "
                            "vector initialization count, training time) (Part 1, Sections "
                            "1.3-1.4 of the notebook).")

h2("1.4 t-SNE visualization of the embedding shift")
p("Figure 1 projects all 5 words' own vectors (before and after fine-tuning) into 2D with "
  "t-SNE, connecting each word's before/after position with an arrow. Figure 2 zooms in on "
  "plot specifically, together with its top-3 neighbors in both spaces.")
add_screenshot_placeholder("Figure 1 - t-SNE plot of all 5 words' before/after positions "
                            "(Part 1, Section 1.7 of the notebook).")
add_screenshot_placeholder("Figure 2 - t-SNE detail plot of 'plot' and its neighbors, before "
                            "vs after fine-tuning (Part 1, Section 1.7 of the notebook).")
p("'plot' visibly moves away from its morphological/literal neighbors (plots, plotting, Plot) "
  "and toward its narrative-sense neighbors (story, storyline, plotline) -- a clean, visual "
  "confirmation of the domain-adaptation effect.")

h2("1.5 Per-word vector drift: original vs. fine-tuned")
add_table([
    ["Word", "Cosine similarity (before vs. after)"],
    ["review", 0.5757],
    ["score", 0.6441],
    ["screen", 0.6652],
    ["plot", 0.6738],
    ["cast", 0.6969],
])
p("Most shifted word: review (cosine similarity = 0.5757). In general news text \"review\" "
  "spans many senses (performance review, code review, review of evidence); IMDB's narrow, "
  "repetitive \"movie review / reader comment\" usage pulls it hardest away from its pretrained "
  "position.")
p("Least shifted word: cast (cosine similarity = 0.6969). \"cast\" already has a strong film/"
  "theater sense in general English (a movie's cast), so IMDB usage reinforces rather than "
  "redefines its pretrained meaning.")
add_screenshot_placeholder("Notebook cell output showing the vector-drift table and the "
                            "most-shifted / least-shifted print statement (Part 1, Section 1.8 "
                            "of the notebook).")

# ============================================================ Part 2
h1("2. Retrieval-Augmented Generation (RAG) Pipeline")

h2("2.1 Setup")
p("Corpus: Wikipedia pages for 10 movies (Inception, The Matrix, Titanic (1997), The Godfather, "
  "Pulp Fiction, Forrest Gump, The Dark Knight, Jurassic Park, The Shawshank Redemption, "
  "Interstellar), fetched with the wikipedia package and loaded with LangChain's DirectoryLoader "
  "+ TextLoader. Splitter: LangChain RecursiveCharacterTextSplitter, chunk_size=500, "
  "chunk_overlap=50, producing 1,970 chunks from the 10 documents. Embeddings: sentence-"
  "transformers/all-MiniLM-L6-v2 via HuggingFaceEmbeddings. Vector store: FAISS, top-k=3 "
  "retriever. LLM: Qwen/Qwen2.5-0.5B-Instruct (local, greedy decoding for reproducibility), "
  "called through transformers.pipeline. The pipeline is built from explicit components -- a "
  "PromptTemplate, the FAISS retriever, and a direct LLM call -- rather than a single pre-built "
  "chain such as RetrievalQA.")
add_screenshot_placeholder("Notebook cell output showing the 10 loaded documents (with char "
                            "counts) and the chunk count after splitting (Part 2, Sections "
                            "2.1-2.3 of the notebook).")

h2("2.2 Five questions, retrieved chunks, and generated answers (original config: 500/50)")
add_table([
    ["#", "Question", "Top retrieved chunk (source)", "Generated answer", "Correct?"],
    [1, "Who directed Inception and who composed its music?", "Inception.txt",
     "Christopher Nolan directed Inception, while Hans Zimmer composed its music.", "Yes"],
    [2, "What is the profession of Andy Dufresne before he was imprisoned in The Shawshank "
        "Redemption?", "The_Shawshank_Redemption.txt",
     "Before being imprisoned in The Shawshank Redemption, Andy Dufresne worked as a banker.", "Yes"],
    [3, "What ship does the film Titanic depict sinking, and in what year did it sink?",
     "Titanic_1997.txt", "The film Titanic depicts the sinking of the RMS Titanic in 1912.", "Yes"],
    [4, "In The Matrix, what is the name of the character played by Keanu Reeves?",
     "The Matrix.txt", "Keanu Reeves as Neo", "Yes"],
    [5, "What dinosaur species famously breaks out of its paddock in Jurassic Park?",
     "Jurassic_Park.txt", "The famous dinosaur species that famously breaks out of its paddock "
     "in Jurassic Park is the Velociraptor.", "No (correct answer: Tyrannosaurus rex)"],
])
add_screenshot_placeholder("Notebook cell output showing all 5 questions with their retrieved "
                            "chunks (top-3, with source filenames) and generated answers, "
                            "original 500/50 configuration (Part 2, Section 2.5 of the "
                            "notebook).")

h2("2.3 Reconfigured chunk size/overlap (1000/100) for 2 of the 5 questions")
p("The vector store was rebuilt with chunk_size=1000, chunk_overlap=100 (1,112 chunks instead "
  "of 1,970), and the Jurassic Park and Titanic questions were re-run.")
add_table([
    ["Question", "Original (500/50)", "Reconfigured (1000/100)"],
    ["Jurassic Park paddock dinosaur",
     "Top-3 chunks miss the T. rex-escape sentence entirely (title, a dinosaur-list header, and "
     "an earlier Velociraptor scene retrieved instead) -> LLM answers \"Velociraptor\" (wrong).",
     "Chunk 1 now contains \"...allows a Tyrannosaurus rex to escape and attack the touring "
     "group\" -> retrieval fixed, but the LLM still answers \"Dilophosaurus\" (wrong; a "
     "different dinosaur named two sentences later in the same chunk)."],
    ["Titanic ship & year",
     "Chunk 1 contains \"...sinking of RMS Titanic in 1912.\" -> correct answer.",
     "Same fact still present in chunk 1 -> same correct answer, unaffected by the "
     "reconfiguration."],
])
p("Larger chunks improved retrieval for the Jurassic Park question (the relevant fact is now "
  "present in the top-3), but did not fix the final answer -- the generation step introduced a "
  "new failure mode once several dinosaur names appeared together in one larger chunk. The "
  "Titanic question was already retrieval-solid at 500/50 and stayed solid at 1000/100.")
add_screenshot_placeholder("Notebook cell output showing the rebuilt vector store's chunk "
                            "count and the re-run retrieved chunks/answers for the Jurassic "
                            "Park and Titanic questions under the 1000/100 configuration "
                            "(Part 2, Section 2.6 of the notebook).")

h2("2.4 Manual retrieval-success assessment (original 500/50 configuration)")
add_table([
    ["Question", "Ground-truth passage", "Found in top-3?", "Rank of first relevant chunk"],
    ["Inception director & composer",
     "\"...directed by Christopher Nolan...\" / \"The score for Inception was composed...by "
     "Hans Zimmer\"", "Yes", 1],
    ["Andy Dufresne's profession",
     "\"...banker Andy Dufresne arrives at Shawshank State Prison...\"", "Yes", 1],
    ["Titanic ship & sinking year",
     "\"...based on accounts of the sinking of RMS Titanic in 1912.\"", "Yes", 1],
    ["Matrix character played by Keanu Reeves",
     "\"Keanu Reeves as Neo: A computer programmer...\"", "Yes", 1],
    ["Jurassic Park paddock-breakout dinosaur",
     "\"...allows a Tyrannosaurus rex to escape and attack the touring group.\"", "No", "N/A"],
])
p("Retrieval Success Rate = 4 / 5 = 80%.")
add_screenshot_placeholder("Notebook cell output showing the manual assessment table and the "
                            "computed Retrieval Success Rate (Part 2, Section 2.7 of the "
                            "notebook).")

h2("2.5 RAG failure analysis")
p("Failure 1 -- correct chunk not retrieved, LLM hallucinates (Jurassic Park, original "
  "config, 500/50). The sentence describing the T. rex's paddock escape falls into a chunk "
  "that never surfaces in the top-3; the retriever instead returns the page title, a bare "
  "\"List\" section header, and an earlier scene where a Velociraptor kills a park worker. "
  "Because \"Velociraptor\" is present in the retrieved context (from that unrelated scene) "
  "while the actual answer is not, the LLM confidently answers \"Velociraptor\" -- the correct "
  "document was effectively not retrieved, and the model filled the gap with a superficially "
  "plausible but wrong entity from the context it did receive.")
p("Failure 2 -- correct context retrieved, but the LLM still answers incorrectly (Jurassic "
  "Park, reconfigured, 1000/100). After widening the chunk size, the correct sentence about the "
  "Tyrannosaurus rex is present in the top-1 retrieved chunk. Despite the right fact being "
  "present, the LLM answers \"Dilophosaurus\" -- a different dinosaur mentioned two sentences "
  "later in the same chunk (Nedry is killed by \"a venom-spitting Dilophosaurus\"). This is a "
  "pure generation failure: the small (0.5B parameter) instruction-tuned model had the right "
  "context but failed to correctly attribute the \"escapes its paddock\" action to the right "
  "entity once multiple dinosaur names competed within the same passage. It shows that fixing "
  "retrieval does not guarantee a correct final answer.")
p("Contributing factor -- ambiguous / overlapping entity names. Both failures share a root "
  "cause: the Jurassic Park plot passage densely packs multiple named dinosaur species into "
  "adjacent sentences (Velociraptor, Tyrannosaurus rex, Dilophosaurus, Brachiosaurus). This is a "
  "specific case of ambiguous-entity confusion -- not movie-title ambiguity, but in-document "
  "entity crowding that confuses both the retriever's chunk selection and the reader's "
  "attribution.")

# ============================================================ Part 3
h1("3. Training-Time Optimization Techniques")

h2("3.1 Shared experimental setup")
p("Model: an 8-layer Transformer-encoder classifier (d_model=384, 6 heads, feedforward "
  "dim=1536). Data: a single fixed synthetic batch (batch_size=64, seq_len=256, embed_dim=384, "
  "random binary labels), generated once with a fixed seed and reused unchanged across every "
  "experiment. Training: 20 steps per run, repeatedly training on the same fixed batch. "
  "Hardware: Apple Silicon GPU (Metal / MPS backend); no CUDA GPU was available, so CUDA-"
  "specific memory APIs (torch.cuda.memory_allocated) were replaced with "
  "torch.mps.current_allocated_memory(), sampled at the point in each training step where the "
  "technique's effect would show up (peak activation memory is highest right after the forward "
  "pass, before backward frees intermediate activations).")

h2("3.2 Tensor Creation: CPU vs. GPU")
p("Tensor creation snippet:")
p("x_cpu = torch.randn(64, 256, 384)                      # lives in system RAM, ops run on CPU\n"
  "x_gpu = torch.randn(64, 256, 384, device=\"mps\")       # created directly on GPU")
add_table([
    ["Device", "Time (20 steps)", "Peak memory", "Loss (start -> end)"],
    ["CPU", "67.63 s", "7783.9 MB (process RSS)", "0.7021 -> 0.6972"],
    ["MPS (GPU)", "32.82 s", "7078.0 MB (MPS allocator)", "0.7004 -> 0.6858"],
])
p("MPS speedup over CPU: 2.06x for the identical model, data, batch size, and step count. Final "
  "loss values match closely, confirming both runs perform the same computation on different "
  "hardware. Peak-memory figures are not directly comparable across devices here: the CPU number "
  "is whole-process resident memory, while the MPS number is the PyTorch MPS allocator's "
  "live-tensor count.")
add_screenshot_placeholder("Notebook cell output showing the CPU vs MPS timing/memory/loss "
                            "print statements and the computed speedup (Part 3, Section 1 of "
                            "the notebook).")

h2("3.3 Weight Initialization")
p("Compared PyTorch's default (Kaiming-uniform) initialization against Xavier/Glorot uniform "
  "and a degenerate all-zeros baseline, applied to every nn.Linear layer.")
add_table([
    ["Init scheme", "Time (20 steps)", "Peak memory", "Loss (start -> end)"],
    ["default (Kaiming)", "42.87 s", "7077.4 MB", "0.7004 -> 0.6858"],
    ["Xavier", "43.19 s", "7079.1 MB", "0.7439 -> 0.7452"],
    ["zeros", "43.60 s", "7078.1 MB", "0.6931 -> 0.3431"],
])
p("Default init starts near ln(2) = 0.693 (expected for a balanced 2-class softmax at "
  "initialization) and descends smoothly. Xavier starts at a much higher initial loss for this "
  "architecture (the Transformer layers' own LayerNorm-scaled sublayers interact with Xavier's "
  "differently-scaled weights) and does not recover within 20 steps -- it plateaus/drifts "
  "slightly upward instead of decreasing. Zeros-initialized linear layers still manage to "
  "reduce loss substantially, because the model's residual (\"skip\") connections let gradients "
  "flow through the untouched input path even when a sublayer's own weights start at exactly "
  "zero -- residual architectures are comparatively robust to some forms of poor "
  "initialization, unlike a plain deep MLP with zero-initialized weights, which would never "
  "break symmetry and would not train at all.")
add_screenshot_placeholder("Notebook cell output showing the three initialization schemes' "
                            "time/memory/loss print statements, and the loss-over-steps line "
                            "plot comparing all three (Part 3, Section 2 of the notebook).")

h2("3.4 Activation Checkpointing")
p("Checkpointing snippet:")
p("x = layer(x)                                              # normal: activations retained\n"
  "x = torch.utils.checkpoint.checkpoint(layer, x, use_reentrant=False)  # recomputed in backward")
add_table([
    ["Configuration", "Time (20 steps)", "Peak activation memory", "Loss (start -> end)"],
    ["No checkpointing", "47.68 s", "7078.5 MB", "0.7004 -> 0.6858"],
    ["With checkpointing", "52.39 s", "357.1 MB", "0.7004 -> 0.6886"],
])
p("Peak activation memory is reduced 19.8x, at the cost of a 1.10x increase in wall-clock time "
  "(every checkpointed layer's forward computation runs a second time during backward). Final "
  "training loss is essentially unchanged -- checkpointing is purely a memory/compute "
  "trade-off, not a modeling change.")
add_screenshot_placeholder("Notebook cell output showing the with/without-checkpointing time, "
                            "peak-activation-memory, and loss print statements, plus the "
                            "computed memory-reduction and time-overhead ratios (Part 3, "
                            "Section 3 of the notebook).")

h2("3.5 Gradient Accumulation")
p("Gradient accumulation snippet:")
p("optimizer.zero_grad()\n"
  "for micro_x, micro_y in micro_batches:                # 4 micro-batches of size 16 = batch 64\n"
  "    out = model(micro_x)\n"
  "    loss = loss_fn(out, micro_y) / len(micro_batches)\n"
  "    loss.backward()                                    # accumulates into .grad\n"
  "optimizer.step()                                       # one update for the whole batch")
add_table([
    ["Configuration", "Time (20 steps)", "Peak memory", "Loss (start -> end)"],
    ["Baseline (batch=64, 1 step)", "37.29 s", "7078.9 MB", "0.7004 -> 0.6858"],
    ["Gradient accumulation (4x micro-batch=16)", "37.04 s", "1971.2 MB", "0.6989 -> 0.6908"],
])
p("Splitting the same 64-example batch into 4 micro-batches of 16 reduces peak activation "
  "memory by roughly 3.6x with essentially the same wall-clock time and a nearly identical loss "
  "trajectory to the full-batch baseline, confirming that gradient accumulation reproduces "
  "full-batch training dynamics while trading a small amount of loop overhead for a large "
  "memory reduction.")
add_screenshot_placeholder("Notebook cell output showing the baseline vs gradient-accumulation "
                            "time/memory/loss print statements (Part 3, Section 4 of the "
                            "notebook).")

h2("3.6 Mixed Precision Training")
p("Mixed precision snippet:")
p("optimizer.zero_grad()\n"
  "with torch.autocast(device_type=\"mps\", dtype=torch.float16):\n"
  "    out = model(x)\n"
  "    loss = loss_fn(out, y)\n"
  "loss.backward()\n"
  "optimizer.step()")
add_table([
    ["Precision", "Time (20 steps)", "Peak memory", "Loss (start -> end)"],
    ["fp32 (baseline)", "37.46 s", "7081.1 MB", "0.7004 -> 0.6858"],
    ["Mixed precision (fp16 autocast)", "35.09 s", "5125.0 MB", "0.7004 -> 0.6853"],
])
p("fp16 autocast reduces peak memory by about 28% and gives a modest (~6%) speedup on this "
  "Apple Silicon GPU, with final loss essentially matching the fp32 baseline. MPS's fp16 "
  "support is comparatively less mature than NVIDIA Tensor Cores under CUDA; on a CUDA GPU, "
  "mixed precision typically yields a much larger speedup (often 2-3x) because Tensor Cores "
  "execute fp16 matmuls natively at higher throughput than fp32, whereas Apple's GPU has no "
  "equivalent dedicated low-precision compute path.")
add_screenshot_placeholder("Notebook cell output showing the fp32 vs mixed-precision time/"
                            "memory/loss print statements (Part 3, Section 5 of the notebook).")

h2("3.7 Summary across all 5 techniques")
add_table([
    ["Technique", "Config A", "Time A (s)", "Mem A (MB)", "Loss A (final)",
     "Config B", "Time B (s)", "Mem B (MB)", "Loss B (final)"],
    ["Tensor creation (GPU vs CPU)", "CPU", 67.63, 7783.9, 0.6972,
     "MPS", 32.82, 7078.0, 0.6858],
    ["Weight init: default", "default", 42.87, 7077.4, 0.6858, "-", "-", "-", "-"],
    ["Weight init: xavier", "xavier", 43.19, 7079.1, 0.7452, "-", "-", "-", "-"],
    ["Weight init: zeros", "zeros", 43.60, 7078.1, 0.3431, "-", "-", "-", "-"],
    ["Activation checkpointing", "no checkpointing", 47.68, 7078.5, 0.6858,
     "with checkpointing", 52.39, 357.1, 0.6886],
    ["Gradient accumulation", "baseline batch=64", 37.29, 7078.9, 0.6858,
     "4x micro-batch=16", 37.04, 1971.2, 0.6908],
    ["Mixed precision", "fp32", 37.46, 7081.1, 0.6858,
     "fp16 autocast", 35.09, 5125.0, 0.6853],
])
add_screenshot_placeholder("Notebook cell output showing the final combined pandas summary "
                            "table across all 5 techniques (Part 3, Summary section of the "
                            "notebook).")

p("Overall takeaways: tensor creation (CPU vs. GPU) is the foundational choice, since everything "
  "downstream runs on whichever device the tensors already live on, and it gave the largest "
  "single speedup of any technique tested. Weight initialization does not change memory or "
  "speed at all -- it only changes the starting point and shape of the loss curve, and residual "
  "architectures are forgiving of even pathological (all-zero) initialization in a way plain "
  "deep feedforward networks would not be. Activation checkpointing and gradient accumulation "
  "both trade a modest time/complexity cost for a large reduction in peak activation memory "
  "without changing what the model learns -- they are complementary techniques for fitting "
  "bigger models/batches into limited GPU memory. Mixed precision reduces memory and gives a "
  "modest speedup here; on CUDA hardware with Tensor Cores, the speedup is typically much "
  "larger.")

# ============================================================ AI Use
h1("4. AI Use")
p("[TODO: fill in your own AI-use disclosure here, in the same format as assignments/"
  "assignment-1/AI_USE .md -- what you used AI for, one thing it got wrong, and how you "
  "checked the results yourself.]")

doc.save(OUT)
print("wrote", OUT)
